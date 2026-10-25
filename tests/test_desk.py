"""Tailor Desk: postings from text, PDFs and links; facts and flags; stages; tailoring; and the page's protections."""
from __future__ import annotations

import http.server
import json
import socket
import sys
import tempfile
import threading
import unittest
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from fastapi.testclient import TestClient  # noqa: E402

from tailor.config import Settings, apply_environment, load_settings  # noqa: E402
from tailor.fill.profile import Profile  # noqa: E402
from tailor.llm import StubClient  # noqa: E402
from tailor.server import desk  # noqa: E402
from tailor.server.app import create_app  # noqa: E402

TOKEN = "desk-token-456"
H = {"X-Tailor-Token": TOKEN}
PROFILE = Profile(first_name="Jane", last_name="Doe", email="jane@example.com", phone="(555) 010-0000")
BASE_PATH = load_settings().profile_dir / "base_resume.json"

SOON = (date.today() + timedelta(days=4)).strftime("%B %d, %Y")
POSTING = f"""Associate Product Manager, University Graduate
Acme Robotics · Mountain View, CA

Please apply by {SOON}.

Minimum qualifications:
Bachelor's degree in Computer Science or a related technical field.
3+ years of experience in product management or software development.
Experience building cross-functional relationships and delivering technical presentations.

Preferred qualifications:
Experience applying AI/ML concepts to build products. Experience with data analysis and user research.

Responsibilities:
Work with engineering, design and research to drive product development from conception to launch.
Use data to prioritize features, define roadmaps and measure outcomes for millions of users.
"""

OTHER_JOBS = "".join(f"<li><h3>Other role {i}: Commercial Program Manager</h3><p>Chicago, IL</p></li>" for i in range(40))
PAGE = f"""<!doctype html><html><head><title>Careers - Acme</title><meta property="og:site_name" content="Acme Robotics"></head>
<body><nav>Home Jobs Students</nav><div><h2>Job search results</h2><ul>{OTHER_JOBS}</ul></div>
<main><div><h2>Associate Product Manager, University Graduate</h2><p>Acme Robotics · Mountain View</p>
<div><h3>Minimum qualifications:</h3><p>{POSTING.split('Minimum qualifications:')[1].replace(chr(10), ' ')}</p>
<p>{'We build helpful robots for everyone. ' * 30}</p></div></div></main></body></html>"""


WORKDAY_PAGE = f"""<!doctype html><html><head><title>Workday</title></head><body>
<div>In order to provide a more relevant experience for you, our Workday Recruiting platform uses cookies. Accept Cookies</div>
<h1>CAREERS AT ACME</h1><nav>Sign In Home Search for Jobs</nav>
<h2 data-automation-id="jobPostingHeader">Strategy Analyst, Early Career</h2>
<div data-automation-id="job-posting-details">locations New York, NY · time type Full time</div>
<div data-automation-id="jobPostingDescription"><p>{POSTING.replace(chr(10), '</p><p>')}</p></div></body></html>"""


class _Site(http.server.BaseHTTPRequestHandler):
    def do_GET(self):
        body = (WORKDAY_PAGE if "myworkday" in self.path else PAGE).encode()
        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, *a):
        pass


def free_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


class DeskTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.settings = Settings(data_dir=Path(self.tmp.name))
        self.settings.ensure_dirs()
        apply_environment(load_settings())  # the browsers live under the real data folder
        self.c = TestClient(create_app(self.settings, StubClient(), PROFILE, TOKEN), base_url="http://127.0.0.1:8765")

    def tearDown(self):
        desk.ALLOW_PRIVATE_HOSTS = False
        self.tmp.cleanup()

    def add(self, **body):
        return self.c.post("/api/desk/postings", json=body, headers=H)

    def test_the_apply_queue_hands_out_postings_one_at_a_time(self):
        body = "Product Manager. Requirements: product strategy, SQL, user research, roadmaps. " * 10
        a = self.add(text=body, url="https://jobs.lever.co/acme/11111111-1111-1111-1111-111111111111", source="text").json()["id"]
        b = self.add(text=body + " B", url="https://job-boards.greenhouse.io/acme/jobs/222", source="text").json()["id"]
        done = self.add(text=body + " C", url="https://job-boards.greenhouse.io/acme/jobs/333", source="text").json()["id"]
        harbor = self.add(text=body + " D", url="https://lakeviewharbor.12twenty.com/jobPostings#/jobPostings/1", source="text").json()["id"]
        self.c.patch(f"/api/desk/postings/{done}", json={"stage": "applied"}, headers=H)
        q = self.c.post("/api/desk/queue", json={"ids": [a, done, harbor, b]}, headers=H).json()
        self.assertEqual(q["ids"], [a, b])  # applied and Harbor Recruit postings are left out
        first = self.c.post("/api/desk/queue/next", headers=H).json()
        self.assertEqual((first["id"], first["left"]), (a, 1))
        self.assertIn("#tailor-fill=", first["url"])
        self.assertTrue(first["url"].endswith("/apply#tailor-fill=" + first["url"].split("#tailor-fill=")[1]))  # Lever's form page
        self.assertEqual(self.c.post("/api/desk/queue/next", headers=H).json()["id"], b)
        self.assertTrue(self.c.post("/api/desk/queue/next", headers=H).json()["done"])
        self.assertEqual(self.c.post("/api/desk/queue/next").status_code, 401)

    # ---- the page and its protections ----
    def test_page_is_served_locally_with_the_token_and_safe_headers(self):
        r = self.c.get("/desk")
        self.assertEqual(r.status_code, 200)
        self.assertIn(f'content="{TOKEN}"', r.text)
        self.assertNotIn("__TAILOR_TOKEN__", r.text)
        self.assertEqual(r.headers["x-frame-options"], "DENY")
        self.assertEqual(r.headers["cache-control"], "no-store")
        foreign = TestClient(self.c.app, base_url="http://evil.example").get("/desk")
        self.assertEqual(foreign.status_code, 421)  # a rebinding page cannot read the token

    def test_every_desk_call_needs_the_token(self):
        for method, path in [("GET", "/api/desk/postings"), ("POST", "/api/desk/postings"),
                             ("DELETE", "/api/desk/postings/abc"), ("POST", "/api/desk/postings/abc/tailor")]:
            self.assertEqual(self.c.request(method, path, json={}).status_code, 401, path)

    # ---- adding postings ----
    def test_text_posting_gets_facts_and_flags(self):
        p = self.add(text=POSTING, url="https://acme.example/jobs/1", source="text").json()
        self.assertEqual(p["title"], "Associate Product Manager, University Graduate")
        self.assertEqual(p["facts"]["min_years"], 3)
        self.assertEqual(p["facts"]["deadline"], (date.today() + timedelta(days=4)).isoformat())
        texts = [f["text"] for f in p["flags"]]
        self.assertIn("3+ years asked", texts)
        self.assertTrue(any(t.startswith("due in 4") for t in texts), texts)
        self.assertNotIn("text", p)  # the list never carries full texts
        again = self.add(text=POSTING, url="https://acme.example/jobs/1").json()
        self.assertTrue(again["duplicate"])
        self.assertEqual(len(self.c.get("/api/desk/postings", headers=H).json()), 1)
        self.assertIn("Minimum qualifications", self.c.get(f"/api/desk/postings/{p['id']}", headers=H).json()["text"])

    def test_pdf_posting(self):
        from playwright.sync_api import sync_playwright

        with sync_playwright() as pw:
            b = pw.chromium.launch()
            page = b.new_page()
            page.set_content("<pre style='font:12px sans-serif'>" + POSTING + "</pre>")
            pdf = page.pdf()
            b.close()
        r = self.c.post("/api/desk/postings/pdf", files=[("files", ("acme_apm.pdf", pdf, "application/pdf"))], headers=H)
        self.assertEqual(r.status_code, 200, r.text)
        p = r.json()[0]
        self.assertEqual((p["source"], p["file"]), ("pdf", "acme_apm.pdf"))
        self.assertIn("Associate Product Manager", p["title"])
        self.assertEqual(p["facts"]["min_years"], 3)
        bad = self.c.post("/api/desk/postings/pdf", files=[("files", ("x.docx", b"PK", "application/octet-stream"))], headers=H)
        self.assertEqual(bad.status_code, 422)

    def test_link_keeps_only_the_posting(self):
        desk.ALLOW_PRIVATE_HOSTS = True
        port = free_port()
        srv = http.server.ThreadingHTTPServer(("127.0.0.1", port), _Site)
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        try:
            p = self.add(url=f"http://127.0.0.1:{port}/jobs/apm").json()
        finally:
            srv.shutdown()
        self.assertEqual(p["title"], "Associate Product Manager, University Graduate")
        self.assertEqual(p["company"], "Acme Robotics")
        text = self.c.get(f"/api/desk/postings/{p['id']}", headers=H).json()["text"]
        self.assertIn("Minimum qualifications", text)
        self.assertNotIn("Commercial Program Manager", text)  # the list of other jobs is left out

    def test_workday_link_keeps_the_posting_not_the_cookie_banner(self):
        desk.ALLOW_PRIVATE_HOSTS = True
        port = free_port()
        srv = http.server.ThreadingHTTPServer(("127.0.0.1", port), _Site)
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        try:
            p = self.add(url=f"http://127.0.0.1:{port}/myworkday/job/1").json()
        finally:
            srv.shutdown()
        self.assertEqual(p["title"], "Strategy Analyst, Early Career")
        text = self.c.get(f"/api/desk/postings/{p['id']}", headers=H).json()["text"]
        self.assertTrue(text.startswith("Strategy Analyst, Early Career"))
        self.assertNotIn("cookies", text)
        self.assertIn("Minimum qualifications", text)

    def test_links_into_this_computer_or_network_are_refused(self):
        for url in ("http://127.0.0.1:8765/desk", "http://localhost/admin", "http://192.168.1.1/", "file:///etc/passwd",
                    "ftp://acme.example/job"):
            r = self.add(url=url)
            self.assertEqual(r.status_code, 422, url)

    def test_nothing_to_add(self):
        self.assertEqual(self.add().status_code, 422)
        self.assertEqual(self.add(text="Apply now.").status_code, 422)

    # ---- stages and removal ----
    def test_stages_and_delete(self):
        pid = self.add(text=POSTING, source="text").json()["id"]
        ok = self.c.patch(f"/api/desk/postings/{pid}", json={"stage": "applied", "notes": "referral from Sam"}, headers=H)
        self.assertEqual((ok.json()["stage"], ok.json()["notes"]), ("applied", "referral from Sam"))
        self.assertEqual(self.c.patch(f"/api/desk/postings/{pid}", json={"stage": "hired"}, headers=H).status_code, 422)
        self.assertEqual(self.c.delete(f"/api/desk/postings/{pid}", headers=H).status_code, 200)
        self.assertEqual(self.c.get("/api/desk/postings", headers=H).json(), [])
        self.assertEqual(self.c.delete(f"/api/desk/postings/{pid}", headers=H).status_code, 404)

    # ---- tailoring ----
    @unittest.skipUnless(BASE_PATH.exists(), "needs the base resume")
    def test_tailor_a_posting_and_fetch_its_resume(self):
        from tailor.resume.schema import Resume

        base = Resume.model_validate_json(BASE_PATH.read_text(encoding="utf-8"))
        c = TestClient(create_app(self.settings, StubClient(), PROFILE, TOKEN, base_resume=base), base_url="http://127.0.0.1:8765")
        pid = c.post("/api/desk/postings", json={"text": POSTING, "source": "text"}, headers=H).json()["id"]
        p = c.post(f"/api/desk/postings/{pid}/tailor", json={"instruction": "lean on product development"}, headers=H).json()
        self.assertEqual(p["status"], "tailored", p["result"])
        self.assertEqual(p["instruction"], "lean on product development")
        rid = p["result"]["id"]
        pdf = c.get("/api/resume", params={"id": rid}, headers=H)
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        # after a restart the Desk's resumes can still be fetched (the file is on disk), but made-up ids cannot
        c2 = TestClient(create_app(self.settings, StubClient(), PROFILE, TOKEN, base_resume=base), base_url="http://127.0.0.1:8765")
        self.assertEqual(c2.get("/api/resume", params={"id": rid}, headers=H).status_code, 200)
        self.assertEqual(c2.get("/api/resume", params={"id": "0123456789ab"}, headers=H).status_code, 404)


class AppliedTests(unittest.TestCase):
    ITEMS = [{"id": "gh", "url": "https://job-boards.greenhouse.io/figma/jobs/6119180004", "title": "Business Operations"},
             {"id": "rb", "url": "https://careers.roblox.com/jobs/8143976?gh_jid=8143976",
              "title": "[2027] Associate Product Manager, Early Career"},
             {"id": "pdf", "url": "", "title": "Digital Assets Consultant"}]

    def m(self, url, title="", text=""):
        hit = desk.match_applied(self.ITEMS, url, title, text)
        return hit["id"] if hit else None

    def test_matches_by_link_job_number_or_title_on_the_same_site(self):
        self.assertEqual(self.m("https://job-boards.greenhouse.io/figma/jobs/6119180004/confirmation"), "gh")
        self.assertEqual(self.m("https://boards.greenhouse.io/embed/job_app/confirmation?token=6119180004"), "gh")
        self.assertEqual(self.m("https://job-boards.greenhouse.io/embed/job_app/confirmation?for=figma&token=6119180004"), "gh")
        self.assertEqual(self.m("https://careers.roblox.com/thank-you", "Roblox Careers",
                                "Thank you for applying to [2027] Associate Product Manager, Early Career"), "rb")

    def test_the_job_number_matches_across_sites(self):
        items = [{"id": "rb", "url": "https://careers.roblox.com/jobs/8143976?gh_jid=8143976", "title": "APM"}]
        hit = desk.match_applied(items, "https://job-boards.greenhouse.io/embed/job_app/confirmation?for=roblox&token=8143976", "", "Thanks")
        self.assertEqual(hit["id"], "rb")

    def test_never_guesses(self):
        self.assertIsNone(self.m("https://jobs.lever.co/acme/thanks", "", "Thank you for applying: Business Operations"))
        self.assertIsNone(self.m("https://careers.roblox.com/thank-you", "", "Thank you for applying!"))  # which role?

    def test_the_tracker_names_the_resume_the_form_got(self):
        from tailor.server.desk import _resume_cell
        self.assertEqual(_resume_cell({"sent_resume": "abcdef123456", "sent_ats": 78}, {}), "tailored abcdef123456.pdf (ATS 78%)")
        self.assertEqual(_resume_cell({"sent_resume": ""}, {"id": "ffffffffffff"}), "default resume")
        self.assertEqual(_resume_cell({}, {"id": "ffffffffffff"}), "tailored ffffffffffff.pdf")  # older entries
        tmp = tempfile.TemporaryDirectory()
        try:
            st = Settings(data_dir=Path(tmp.name))
            st.ensure_dirs()
            (st.output_dir / "tailored").mkdir(parents=True, exist_ok=True)
            (st.output_dir / "tailored" / "0123456789ab.json").write_text('{"ats_after": 81}', encoding="utf-8")
            c = TestClient(create_app(st, StubClient(), PROFILE, TOKEN), base_url="http://127.0.0.1:8765")
            pid = c.post("/api/desk/postings", headers=H, json={"text": POSTING, "source": "text",
                                                                 "url": "https://job-boards.greenhouse.io/acme/jobs/7654321"}).json()["id"]
            c.post("/api/desk/applied", headers=H, json={"url": "https://job-boards.greenhouse.io/acme/jobs/7654321/confirmation",
                                                         "resume_id": "0123456789ab"})
            from openpyxl import load_workbook
            row = [c.value for c in load_workbook(st.data_dir / "Application_Tracker.xlsx").active[2]]
            self.assertEqual(row[7], "tailored 0123456789ab.pdf (ATS 81%)")
            self.assertEqual(c.post("/api/desk/applied", headers=H, json={"url": "x", "resume_id": "../etc"}).status_code, 422)
            self.assertTrue(pid)
        finally:
            tmp.cleanup()

    def test_endpoint_marks_applied_once(self):
        tmp = tempfile.TemporaryDirectory()
        try:
            st = Settings(data_dir=Path(tmp.name))
            st.ensure_dirs()
            c = TestClient(create_app(st, StubClient(), PROFILE, TOKEN), base_url="http://127.0.0.1:8765")
            pid = c.post("/api/desk/postings", headers=H, json={"text": POSTING, "source": "text",
                                                                 "url": "https://job-boards.greenhouse.io/acme/jobs/1234567"}).json()["id"]
            r = c.post("/api/desk/applied", headers=H, json={"url": "https://job-boards.greenhouse.io/acme/jobs/1234567/confirmation"}).json()
            self.assertEqual((r["matched"], r["already"]), (pid, False))
            self.assertEqual(c.get(f"/api/desk/postings/{pid}", headers=H).json()["stage"], "applied")
            self.assertTrue(c.post("/api/desk/applied", headers=H, json={"url": "https://job-boards.greenhouse.io/acme/jobs/1234567/x"}).json()["already"])
            self.assertIsNone(c.post("/api/desk/applied", headers=H, json={"url": "https://other.example/thanks"}).json()["matched"])
            self.assertEqual(c.post("/api/desk/applied", json={}).status_code, 401)
            # the Excel tracker: rewritten when something is marked applied
            from datetime import date as _d
            from openpyxl import load_workbook
            ws = load_workbook(st.data_dir / "Application_Tracker.xlsx").active
            self.assertEqual([c.value for c in ws[1]][:6], ["Date applied", "Company", "Job title", "Link", "Qualifications", "Description"])
            row = [c.value for c in ws[2]]
            self.assertEqual(row[0], _d.today().isoformat())
            self.assertEqual(row[2], "Associate Product Manager, University Graduate")
            self.assertEqual(row[3], "https://job-boards.greenhouse.io/acme/jobs/1234567")
            self.assertIn("cross-functional", row[4])
            self.assertIn("product development", row[5])
            self.assertEqual(ws.cell(row=2, column=4).hyperlink.target, row[3])
            self.assertEqual(c.get("/api/desk/tracker", headers=H).status_code, 200)

            # answers sent with the confirmation go to the library, except sensitive ones
            c.post("/api/desk/applied", headers=H, json={"url": "https://job-boards.greenhouse.io/acme/jobs/1234567/done", "answers": [
                {"question": "Why do you want to join Acme?", "text": "I like turning messy questions into decisions, as I did at Lakeside Health."},
                {"question": "Are you legally authorized to work in the United States?", "text": "Yes, and this should never be kept by Tailor."}]})
            lib = json.loads((st.profile_dir / "answer_library.json").read_text())
            self.assertEqual([e["question"] for e in lib], ["Why do you want to join Acme?"])
            self.assertEqual(lib[0]["company"], "Acme")
        finally:
            tmp.cleanup()


class PdfHeaderTests(unittest.TestCase):
    """Printed career pages: icon-font characters and the site's menu come before the title."""

    def test_title_skips_icons_and_menus_and_joins_a_split_title(self):
        accenture = ("\ue80e        \ue82d   \ue82c  USA   \ue80a\nCareers      Job Search     Saved Jobs    Access Application\n"
                     "Digital Assets Consultant\nMgmt Consulting Team Lead/Consultant  |  Mid-Level  |  Full time\n")
        self.assertEqual(desk._title_from_text(accenture), "Digital Assets Consultant")
        pwc = "Strategy& Deals Strategy Senior\nAssociate - Summer/Fall 2027\nAvailable in 3 locations See all\n"
        self.assertEqual(desk._title_from_text(pwc), "Strategy& Deals Strategy Senior Associate - Summer/Fall 2027")
        self.assertEqual(desk._title_from_text("\ue80e\n\n"), "Untitled posting")

    def test_company_from_the_file_name_only_when_the_posting_says_it(self):
        self.assertEqual(desk.company_from_file("pwc_strategyand_deals.pdf", "Join PwC Strategy& as a Senior Associate"), "PwC")
        self.assertEqual(desk.company_from_file("accenture-consultant.pdf", "Accenture is growing"), "Accenture")
        self.assertEqual(desk.company_from_file("posting (3).pdf", "Acme is hiring"), "")
        self.assertEqual(desk.company_from_file("my_resume_target.pdf", "Acme is hiring"), "")


class CompanyGuessTests(unittest.TestCase):
    def test_from_job_board_links(self):
        self.assertEqual(desk.company_from("https://job-boards.greenhouse.io/figma/jobs/123"), "Figma")
        self.assertEqual(desk.company_from("https://jobs.lever.co/acme-labs/abc"), "Acme Labs")
        self.assertEqual(desk.company_from("https://kpmg.wd5.myworkdayjobs.com/en-US/x"), "Kpmg")
        self.assertEqual(desk.company_from("https://www.google.com/about/careers/applications/jobs/results/1"), "Google")
        self.assertEqual(desk.company_from("https://careers.roblox.com/jobs/1"), "Roblox")
        self.assertEqual(desk.company_from("https://lifeattiktok.com/search/7667791791502821637"), "TikTok")
        self.assertEqual(desk.company_from("https://x.example", site="Menlo Innovations"), "Menlo Innovations")
        self.assertEqual(desk.company_from("https://www.linkedin.com/jobs/1", site="LinkedIn"), "Linkedin")


if __name__ == "__main__":
    unittest.main()


class GreenhouseFormTests(unittest.TestCase):
    """2026-10-02: Stripe, Roblox and Databricks postings live on their own sites; 🚀 Apply opens Greenhouse's form."""

    def test_company_site_posting_opens_the_greenhouse_form(self):
        from types import SimpleNamespace
        from unittest import mock

        from tailor.server import desk

        def fake_get(url, **kw):
            ok = "for=roblox&" in url
            return SimpleNamespace(status_code=200, url=url if ok else "https://job-boards.greenhouse.io/embed/job_board?error=true",
                                   text="<title>Job Application for APM at Roblox</title>" if ok else "Jobs")

        with tempfile.TemporaryDirectory() as tmp, mock.patch.object(desk.httpx, "get", fake_get):
            settings = SimpleNamespace(data_dir=Path(tmp))
            self.assertEqual(desk.greenhouse_form("https://careers.roblox.com/jobs/8143976?gh_jid=8143976", settings),
                             "https://job-boards.greenhouse.io/embed/job_app?for=roblox&token=8143976")
            self.assertIsNone(desk.greenhouse_form("https://example.com/jobs?gh_jid=1", settings))
            self.assertIsNone(desk.greenhouse_form("https://job-boards.greenhouse.io/oscar/jobs/8132502", settings))


class ExtensionDownloadTests(unittest.TestCase):
    def test_the_zip_holds_the_extension_and_no_token(self):
        import io
        import zipfile

        from tailor.server.desk import EXTENSION_DIR

        self.assertTrue((EXTENSION_DIR / "manifest.json").exists())
        names = [p.name for p in EXTENSION_DIR.rglob("*") if p.is_file()]
        for f in ("manifest.json", "content.js", "background.js", "popup.html", "options.html"):
            self.assertIn(f, names)
        self.assertFalse(any("token" in n.lower() for n in names))


class CompanyFromMoreSitesTests(unittest.TestCase):
    def test_smartrecruiters_and_cornerstone_links(self):
        from tailor.server.desk import company_from
        self.assertEqual(company_from("https://jobs.smartrecruiters.com/RolandBerger/744000149023210-consultant"), "Roland Berger")
        self.assertEqual(company_from("https://simon-kucher.csod.com/ux/ats/careersite/6/home/requisition/4397"), "Simon Kucher")


class WordingsTests(unittest.TestCase):
    """The Desk's Wordings tab: proposed wordings are used only after the person approves them."""

    def test_list_approve_reject(self):
        tmp = tempfile.TemporaryDirectory()
        try:
            st = Settings(data_dir=Path(tmp.name))
            st.ensure_dirs()
            orig = "Generated more than $310,000 in annual revenue by uncovering customer needs"
            good = "Uncovered customer needs, generating more than $310,000 in annual revenue"
            bad = "Generated more than $500,000 in annual revenue by uncovering customer needs"
            (st.profile_dir / "phrase_bank.json").write_text(json.dumps({"bullets": [{"entry": "ACME", "role": "Sales", "original": orig,
                "variants": [{"angle": "a", "text": good, "status": "proposed"}, {"angle": "b", "text": bad, "status": "proposed"}]}]}),
                encoding="utf-8")
            c = TestClient(create_app(st, StubClient(), PROFILE, TOKEN), base_url="http://127.0.0.1:8765")
            got = c.get("/api/bank", headers=H).json()
            self.assertEqual([v["status"] for v in got["bullets"][0]["variants"]], ["proposed", "proposed"])
            self.assertEqual(c.post("/api/bank/decide", headers=H, json={"original": orig, "text": good, "status": "approved"}).status_code, 200)
            r = c.post("/api/bank/decide", headers=H, json={"original": orig, "text": bad, "status": "approved"})
            self.assertEqual(r.status_code, 422)  # a changed number never becomes approved
            self.assertEqual(c.post("/api/bank/decide", headers=H, json={"original": orig, "text": bad, "status": "rejected"}).status_code, 200)
            self.assertEqual(c.post("/api/bank/decide", headers=H, json={"original": orig, "text": "nope", "status": "approved"}).status_code, 404)
            self.assertEqual(c.get("/api/bank").status_code, 401)
            saved = json.loads((st.profile_dir / "phrase_bank.json").read_text(encoding="utf-8"))
            self.assertEqual([v["status"] for v in saved["bullets"][0]["variants"]], ["approved", "rejected"])
        finally:
            tmp.cleanup()


class CompanyNameTests(unittest.TestCase):
    """2026-10-14: "Openai" on the Desk; a board's slug gets the proper name, or the posting's own mixed-case spelling."""

    def test_proper_names(self):
        from tailor.server.desk import company_from

        self.assertEqual(company_from("https://jobs.ashbyhq.com/openai/5ce0", "", "Join OpenAI as a strategist"), "OpenAI")
        self.assertEqual(company_from("https://job-boards.greenhouse.io/doordashusa/jobs/1"), "DoorDash")
        self.assertEqual(company_from("https://jobs.lever.co/chartis/1", "", "CHARTIS is hiring. Chartis consultants"), "Chartis")

